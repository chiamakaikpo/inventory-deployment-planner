"""
Independent check of the Excel planner.

Re-implements business rules BR-01 to BR-04 in plain Python from the workbook's INPUT
cells, then compares the results with the values Excel/LibreOffice calculated.
If someone edits a formula and breaks a rule, this test fails.

Run:  python -m pytest tests/        (or: python tests/test_planner.py)
The workbook must have been recalculated first (open and save it in Excel, or run
scripts/recalc.sh) so that formula results are stored in the file.
"""
from pathlib import Path
from math import ceil, floor, isclose
from openpyxl import load_workbook

XLSX = Path(__file__).resolve().parents[1] / "planner" / "Deployment_Planner.xlsx"
wb = load_workbook(XLSX, data_only=True)


def read_inputs():
    wi = wb["Inputs"]
    bands = {wi.cell(r, 1).value: (wi.cell(r, 2).value, wi.cell(r, 3).value, wi.cell(r, 4).value) for r in range(5, 8)}
    floor_days, transfer_lt = wi["B11"].value, wi["B12"].value
    priority = {wi.cell(r, 1).value: wi.cell(r, 2).value == "Y" for r in range(17, 22)}
    skus = {wi.cell(r, 1).value: (wi.cell(r, 2).value, wi.cell(r, 3).value) for r in range(26, 32)}
    ws = wb["Stock"]
    rows = []
    for r in range(5, ws.max_row + 1):
        if ws.cell(r, 1).value:
            rows.append(dict(sku=ws.cell(r, 1).value, depot=ws.cell(r, 2).value, fc=ws.cell(r, 3).value,
                             pos=ws.cell(r, 4).value + ws.cell(r, 5).value))
    return bands, floor_days, transfer_lt, priority, skus, rows


def plan():
    bands, floor_days, transfer_lt, priority, skus, rows = read_inputs()
    for row in rows:
        cls, supply = skus[row["sku"]]
        mn, tgt, mx = bands[cls]                                            # BR-01
        row.update(cls=cls, supply=supply, min=mn, max=mx, prio=priority[row["depot"]],
                   need=max(0, round(tgt * row["fc"] - row["pos"])))
    for sku in skus:
        grp = [r for r in rows if r["sku"] == sku]
        supply = grp[0]["supply"]
        total_need = sum(r["need"] for r in grp)
        if total_need <= supply:
            for r in grp:
                r["alloc"] = r["need"]
            continue
        # BR-02: priority floor first (scaled down if floors alone exceed supply)
        for r in grp:
            r["floor"] = min(r["need"], floor_days * r["fc"]) if r["prio"] else 0
        floors = sum(r["floor"] for r in grp)
        for r in grp:
            r["floor"] = r["floor"] * supply / floors if floors > supply else r["floor"]
            r["rem"] = r["need"] - r["floor"]
        # then the rest in proportion to forecast, each depot capped at its need ("water-filling"):
        # repeatedly share what is left among depots that still need stock
        left = max(0, supply - min(floors, supply))
        share = {id(r): 0.0 for r in grp}
        open_ = [r for r in grp if r["rem"] > 0]
        while left > 1e-9 and open_:
            w = sum(r["fc"] for r in open_)
            lam = left / w
            capped = [r for r in open_ if r["rem"] - share[id(r)] <= lam * r["fc"] + 1e-12]
            if not capped:
                for r in open_:
                    share[id(r)] += lam * r["fc"]
                left = 0
                break
            for r in capped:
                left -= r["rem"] - share[id(r)]
                share[id(r)] = r["rem"]
            open_ = [r for r in open_ if r not in capped]
        for r in grp:
            r["alloc"] = floor(r["floor"] + share[id(r)] + 1e-9)
    for r in rows:                                                          # BR-04
        r["cover"] = (r["pos"] + r["alloc"]) / r["fc"]
        r["status"] = "Below band" if r["cover"] < r["min"] else "Above band" if r["cover"] > r["max"] else "Within band"
    return rows, skus, transfer_lt


def excel_allocation():
    wa = wb["Allocation"]
    return [dict(sku=wa.cell(r, 1).value, depot=wa.cell(r, 2).value, alloc=wa.cell(r, 21).value,
                 cover=wa.cell(r, 22).value, status=wa.cell(r, 25).value)
            for r in range(5, wa.max_row + 1) if wa.cell(r, 1).value]


def test_allocation_matches_rules():
    expected, _, _ = plan()
    got = excel_allocation()
    assert len(got) == len(expected) == 30
    for e, g in zip(expected, got):
        assert (e["sku"], e["depot"]) == (g["sku"], g["depot"])
        assert e["alloc"] == g["alloc"], f"{e['sku']} @ {e['depot']}: python {e['alloc']} vs excel {g['alloc']}"
        assert isclose(e["cover"], g["cover"], abs_tol=1e-6)
        assert e["status"] == g["status"]


def test_never_allocates_more_than_supply():
    rows, skus, _ = plan()
    for sku, (_, supply) in skus.items():
        assert sum(r["alloc"] for r in rows if r["sku"] == sku) <= supply


def test_no_depot_gets_more_than_it_needs():
    rows, _, _ = plan()
    assert all(0 <= r["alloc"] <= r["need"] for r in rows)


def test_short_skus_use_all_supply():
    """When a SKU is short, stock left undeployed should be at most rounding (1 case per depot)."""
    rows, skus, _ = plan()
    for sku, (_, supply) in skus.items():
        grp = [r for r in rows if r["sku"] == sku]
        if sum(r["need"] for r in grp) > supply:
            assert supply - sum(r["alloc"] for r in grp) <= len(grp)


def test_transfer_rule():
    """BR-03: transfer only if the biggest excess covers the biggest shortfall and arrives in time."""
    rows, skus, lt = plan()
    wt = wb["Transfers"]
    for i, sku in enumerate(skus, 5):
        grp = [r for r in rows if r["sku"] == sku]
        for r in grp:
            r["excess"] = max(0, floor(round((r["cover"] - r["max"]) * r["fc"], 6)))
            r["short"] = max(0, ceil(round((r["min"] - r["cover"]) * r["fc"], 6)))
        ex = max(grp, key=lambda r: r["excess"])
        sh = max(grp, key=lambda r: r["short"])
        ok = ex["excess"] > 0 and sh["short"] > 0 and ex["excess"] >= sh["short"] and lt <= sh["cover"]
        assert (wt.cell(i, 8).value == "Transfer") == ok, sku
        assert wt.cell(i, 9).value == (sh["short"] if ok else 0), sku


if __name__ == "__main__":
    for name, fn in list(globals().items()):
        if name.startswith("test_"):
            fn()
            print("passed", name)
