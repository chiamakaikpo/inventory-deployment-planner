"""
Build planner/Deployment_Planner.xlsx: a rule-based inventory deployment planner.

Implements the business rules from the case study:
  BR-01  Each depot has a target days-of-cover band per SKU class (fast / medium / slow).
  BR-02  When supply is short, allocate in proportion to forecast demand, with a minimum
         floor for priority depots.
  BR-03  Suggest an inter-depot transfer only when projected excess at one depot covers a
         projected shortfall at another within the transfer lead time.
  BR-04  A depot/SKU is an exception when projected cover falls outside its band.

All inputs are SYNTHETIC (fictional brewer, anonymised depots and SKUs).
Every calculated cell is an Excel formula, so planners can change inputs and see results.
"""
from pathlib import Path
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter as L

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "planner" / "Deployment_Planner.xlsx"
rng = np.random.default_rng(7)

A = "Arial"
F = Font(name=A, size=10)
FB = Font(name=A, size=10, bold=True)
FIN = Font(name=A, size=10, color="0000FF")              # input cells
FH = Font(name=A, size=10, bold=True, color="FFFFFF")
FT = Font(name=A, size=14, bold=True)
FN = Font(name=A, size=9, italic=True, color="5B6878")
HEAD = PatternFill("solid", fgColor="1F3A68")
INFILL = PatternFill("solid", fgColor="FFF8E1")
S = Side(style="thin", color="C4CCD6")
BOX = Border(left=S, right=S, top=S, bottom=S)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

DEPOTS = [("Depot North", "N"), ("Depot Central", "Y"), ("Depot South-West", "Y"), ("Depot South-East", "N"), ("Depot Coastal", "N")]
DEPOT_SCALE = [0.7, 1.0, 1.4, 0.9, 0.6]
SKUS = [  # name, class, base demand cases/day at scale 1, supply factor vs need (<1 = short)
    ("Lager 60cl bottle", "Fast", 900, 0.80),
    ("Lager 33cl can", "Fast", 650, 1.30),
    ("Malt drink 33cl can", "Medium", 380, 0.70),
    ("Stout 60cl bottle", "Medium", 260, 1.10),
    ("Premium lager 33cl bottle", "Slow", 120, 0.90),
    ("Shandy 33cl can", "Slow", 70, 1.50),
]
BANDS = [("Fast", 5, 8, 12), ("Medium", 7, 12, 18), ("Slow", 10, 18, 28)]

wb = Workbook()


def head(ws, row, headers, col=1):
    for j, h in enumerate(headers):
        c = ws.cell(row=row, column=col + j, value=h)
        c.font, c.fill, c.alignment, c.border = FH, HEAD, CENTER, BOX


def widths(ws, ws_widths):
    for i, w in enumerate(ws_widths, 1):
        ws.column_dimensions[L(i)].width = w


# =========================================================== Read me
ws = wb.active
ws.title = "Read me"
rows = [
    ("Inventory Deployment Planner", FT),
    ("Portfolio project by Chiamaka Ikpo · synthetic data for a fictional national brewer (5 depots, 6 SKUs)", FN),
    ("", F),
    ("What it does", FB),
    ("Takes brewery stock available for deployment and allocates it to depots using agreed rules, then flags exceptions and suggests transfers.", F),
    ("It replaces judgement-based allocation in spreadsheets with rules everyone has agreed, so planners only review exceptions.", F),
    ("", F),
    ("Business rules", FB),
    ("BR-01  Each depot has a target days-of-cover band per SKU class (fast / medium / slow movers).  → Inputs: cover bands", F),
    ("BR-02  When supply is short, allocate in proportion to forecast demand, with a minimum floor for priority depots.  → Allocation tab", F),
    ("BR-03  Suggest an inter-depot transfer only when projected excess at one depot covers a projected shortfall at another within the transfer lead time.  → Transfers tab", F),
    ("BR-04  A depot/SKU is flagged as an exception when projected cover falls outside its band.  → Allocation tab, Status column", F),
    ("", F),
    ("How to use", FB),
    ("1. Update blue cells on Inputs (bands, supply, priority depots) and Stock (forecast, stock on hand, in transit).", F),
    ("2. Review Summary: shortages, exceptions and imbalance before vs after deployment.", F),
    ("3. Review only the exception rows on Allocation and the suggested moves on Transfers.", F),
    ("", F),
    ("Colour key: blue text on cream = input you can change · black = formula, do not overwrite.", FN),
    ("Days of cover = (stock on hand + in transit) ÷ forecast daily demand. Units are cases.", FN),
]
for i, (t, f) in enumerate(rows, 1):
    ws.cell(row=i, column=1, value=t).font = f
ws.column_dimensions["A"].width = 140

# =========================================================== Inputs
wi = wb.create_sheet("Inputs")
wi["A1"], wi["A1"].font = "Inputs", FT
wi["A3"], wi["A3"].font = "Cover bands by SKU class (BR-01), days", FB
head(wi, 4, ["Class", "Min cover", "Target cover", "Max cover"])
for i, b in enumerate(BANDS, 5):
    for j, v in enumerate(b, 1):
        c = wi.cell(row=i, column=j, value=v)
        c.border = BOX
        c.font = FIN if j > 1 else F
        if j > 1:
            c.fill = INFILL
BAND = "Inputs!$A$5:$D$7"

wi["A10"], wi["A10"].font = "Parameters", FB
params = [("Priority floor (days of forecast) – BR-02", 3, "B11"), ("Inter-depot transfer lead time (days) – BR-03", 2, "B12")]
for i, (label, v, ref) in enumerate(params, 11):
    wi.cell(row=i, column=1, value=label).font = F
    c = wi.cell(row=i, column=2, value=v)
    c.font, c.fill, c.border = FIN, INFILL, BOX
FLOOR, TLT = "Inputs!$B$11", "Inputs!$B$12"

wi["A15"], wi["A15"].font = "Depots", FB
head(wi, 16, ["Depot", "Priority depot? (Y/N)"])
for i, (d, p) in enumerate(DEPOTS, 17):
    wi.cell(row=i, column=1, value=d).font = F
    c = wi.cell(row=i, column=2, value=p)
    c.font, c.fill, c.border = FIN, INFILL, BOX
dv = DataValidation(type="list", formula1='"Y,N"')
wi.add_data_validation(dv)
dv.add("B17:B21")
DEP = "Inputs!$A$17:$B$21"
wi["C17"] = "Priority depots serve key accounts"
wi["C17"].font = FN

wi["A24"], wi["A24"].font = "Brewery stock available for deployment this cycle", FB
head(wi, 25, ["SKU", "Class", "Available (cases)"])

# synthetic stock positions first (supply is sized from the need they create)
stock_rows = []
for s_name, s_cls, base, _ in SKUS:
    tgt = next(b[2] for b in BANDS if b[0] == s_cls)
    for (d, _), sc in zip(DEPOTS, DEPOT_SCALE):
        fc = round(base * sc * rng.uniform(0.85, 1.15))
        cover = rng.uniform(0.25, 0.9) * tgt
        if s_name == "Stout 60cl bottle" and d == "Depot North":
            cover = 2.4 * tgt            # one overstocked depot -> transfer candidate
        if s_name == "Premium lager 33cl bottle" and d == "Depot Coastal":
            cover = 2.0 * tgt
        if s_name == "Malt drink 33cl can" and d == "Depot Coastal":
            cover = 2.0 * tgt            # overstocked while the SKU is short -> transfer
        soh = round(fc * cover * 0.75)
        transit = round(fc * cover * 0.25)
        stock_rows.append((s_name, d, fc, soh, transit))

for i, (s_name, s_cls, base, factor) in enumerate(SKUS, 26):
    tgt = next(b[2] for b in BANDS if b[0] == s_cls)
    need = sum(max(0, tgt * fc - (soh + tr)) for n, d, fc, soh, tr in stock_rows if n == s_name)
    supply = int(round(need * factor, -2))
    wi.cell(row=i, column=1, value=s_name).font = F
    c = wi.cell(row=i, column=2, value=s_cls)
    c.font, c.fill, c.border = FIN, INFILL, BOX
    c = wi.cell(row=i, column=3, value=supply)
    c.font, c.fill, c.border, c.number_format = FIN, INFILL, BOX, "#,##0"
dvc = DataValidation(type="list", formula1='"Fast,Medium,Slow"')
wi.add_data_validation(dvc)
dvc.add("B26:B31")
SUP = "Inputs!$A$26:$C$31"
widths(wi, [46, 20, 18, 14])

# =========================================================== Stock (inputs per SKU x depot)
wst = wb.create_sheet("Stock")
wst["A1"], wst["A1"].font = "Stock position and forecast by SKU and depot", FT
wst["A2"], wst["A2"].font = "Inputs: from the SAP stock and forecast extract at the agreed data cut-off (to-be step 1: Sync).", FN
head(wst, 4, ["SKU", "Depot", "Forecast (cases/day)", "Stock on hand", "In transit"])
for i, r in enumerate(stock_rows, 5):
    for j, v in enumerate(r, 1):
        c = wst.cell(row=i, column=j, value=v)
        c.border = BOX
        c.font = F if j <= 2 else FIN
        if j > 2:
            c.fill, c.number_format = INFILL, "#,##0"
N0, N1 = 5, 4 + len(stock_rows)
widths(wst, [28, 20, 18, 15, 13])
wst.freeze_panes = "A5"

# =========================================================== Allocation
wa = wb.create_sheet("Allocation")
wa["A1"], wa["A1"].font = "Allocation (BR-02) and exception status (BR-04)", FT
wa["A2"], wa["A2"].font = "All formulas. Filter Status ≠ 'Within band' to see only the exceptions planners need to review.", FN
cols = ["SKU", "Depot", "Class", "Priority", "Forecast /day", "Position (SOH + transit)", "Cover now (days)",
        "Target cover", "Need to target", "SKU supply", "SKU total need", "Short?", "Priority floor",
        "SKU floor total", "Floor (scaled)", "Remaining need", "Supply after floors", "Need ÷ forecast",
        "Use at this ratio", "Ratio that fits", "Allocation", "Projected cover", "Min band", "Max band", "Status",
        "Status before", "Excess above max", "Shortfall below min", "Key excess", "Key shortfall",
        "Base ratio", "Use at base", "Forecast above base", "Share ratio (λ)", "Fair share"]
head(wa, 4, cols)
R0 = 5
for k in range(len(stock_rows)):
    r, s = R0 + k, N0 + k
    rng_A = f"$A${R0}:$A${R0 + len(stock_rows) - 1}"
    col = lambda c: f"${c}${R0}:${c}${R0 + len(stock_rows) - 1}"
    f = {
        "A": f"=Stock!A{s}",
        "B": f"=Stock!B{s}",
        "C": f"=INDEX(Inputs!$B$26:$B$31,MATCH(A{r},Inputs!$A$26:$A$31,0))",
        "D": f"=INDEX(Inputs!$B$17:$B$21,MATCH(B{r},Inputs!$A$17:$A$21,0))",
        "E": f"=Stock!C{s}",
        "F": f"=Stock!D{s}+Stock!E{s}",
        "G": f"=IFERROR(F{r}/E{r},0)",
        "H": f"=INDEX(Inputs!$C$5:$C$7,MATCH(C{r},Inputs!$A$5:$A$7,0))",
        "I": f"=MAX(0,ROUND(H{r}*E{r}-F{r},0))",
        "J": f"=INDEX(Inputs!$C$26:$C$31,MATCH(A{r},Inputs!$A$26:$A$31,0))",
        "K": f"=SUMIF({rng_A},A{r},{col('I')})",
        "L": f'=IF(K{r}>J{r},"Yes","No")',
        "M": f'=IF(AND(L{r}="Yes",D{r}="Y"),MIN(I{r},{FLOOR}*E{r}),0)',
        "N": f"=SUMIF({rng_A},A{r},{col('M')})",
        "O": f"=IF(N{r}>J{r},M{r}*J{r}/N{r},M{r})",
        "P": f"=I{r}-O{r}",
        "Q": f"=MAX(0,J{r}-MIN(N{r},J{r}))",
        # BR-02 fair share: give each depot  min(remaining need, lambda x forecast)  and solve for the
        # lambda that uses all remaining supply ("water-filling"). R-T find the largest ratio that fits.
        "R": f"=IF(P{r}>0,P{r}/E{r},0)",
        "S": f"=SUMPRODUCT(({rng_A}=A{r})*(({col('P')}<R{r}*{col('E')})*{col('P')}+({col('P')}>=R{r}*{col('E')})*R{r}*{col('E')}))",
        "T": f"=IF(AND(P{r}>0,S{r}<=Q{r}),R{r},0)",
        "U": f'=IF(L{r}="No",I{r},ROUNDDOWN(O{r}+AI{r},0))',
        "V": f"=IFERROR((F{r}+U{r})/E{r},0)",
        "W": f"=INDEX(Inputs!$B$5:$B$7,MATCH(C{r},Inputs!$A$5:$A$7,0))",
        "X": f"=INDEX(Inputs!$D$5:$D$7,MATCH(C{r},Inputs!$A$5:$A$7,0))",
        "Y": f'=IF(V{r}<W{r},"Below band",IF(V{r}>X{r},"Above band","Within band"))',
        "Z": f'=IF(G{r}<W{r},"Below band",IF(G{r}>X{r},"Above band","Within band"))',
        "AA": f"=MAX(0,ROUNDDOWN((V{r}-X{r})*E{r},0))",
        "AB": f"=MAX(0,ROUNDUP((W{r}-V{r})*E{r},0))",
        "AC": f'=A{r}&"|"&AA{r}',
        "AD": f'=A{r}&"|"&AB{r}',
        "AE": f"=_xlfn.MAXIFS({col('T')},{rng_A},A{r})",
        "AF": f"=SUMPRODUCT(({rng_A}=A{r})*(({col('P')}<AE{r}*{col('E')})*{col('P')}+({col('P')}>=AE{r}*{col('E')})*AE{r}*{col('E')}))",
        "AG": f"=SUMIFS({col('E')},{rng_A},A{r},{col('R')},\">\"&AE{r})",
        "AH": f"=IF(AG{r}=0,AE{r},AE{r}+(Q{r}-AF{r})/AG{r})",
        "AI": f"=MIN(P{r},AH{r}*E{r})",
    }
    for c_letter, formula in f.items():
        cell = wa[f"{c_letter}{r}"]
        cell.value, cell.font, cell.border = formula, F, BOX
        if c_letter in ("G", "V", "O", "AI", "AF"):
            cell.number_format = "0.0"
        elif c_letter in ("R", "T", "AE", "AH"):
            cell.number_format = "0.000"
        elif False:
            cell.number_format = "0.0"
        elif c_letter not in ("A", "B", "C", "D", "L", "Y", "Z", "AC", "AD"):
            cell.number_format = "#,##0"
AEND = R0 + len(stock_rows) - 1
for c_letter, w in zip([L(i) for i in range(1, 36)],
                       [26, 18, 9, 8, 10, 12, 10, 9, 10, 10, 10, 7, 9, 9, 9, 10, 10, 9, 10, 9, 11, 10, 8, 8, 13, 13, 10, 11, 20, 20, 10, 10, 9, 10, 9]):
    wa.column_dimensions[c_letter].width = w
wa.row_dimensions[4].height = 42
wa.freeze_panes = "C5"
wa.auto_filter.ref = f"A4:AI{AEND}"
for c_letter in ("Y", "Z"):
    rngs = f"{c_letter}{R0}:{c_letter}{AEND}"
    wa.conditional_formatting.add(rngs, CellIsRule(operator="equal", formula=['"Below band"'], fill=PatternFill("solid", fgColor="FCE6E4"), font=Font(name=A, color="B42318", bold=True)))
    wa.conditional_formatting.add(rngs, CellIsRule(operator="equal", formula=['"Above band"'], fill=PatternFill("solid", fgColor="FBF0D5"), font=Font(name=A, color="9A6300", bold=True)))
    wa.conditional_formatting.add(rngs, CellIsRule(operator="equal", formula=['"Within band"'], fill=PatternFill("solid", fgColor="E2F3EA"), font=Font(name=A, color="1F7A4D")))
wa.conditional_formatting.add(f"U{R0}:U{AEND}", FormulaRule(formula=[f"$L{R0}=\"Yes\""], fill=PatternFill("solid", fgColor="E7EDFA")))
# working columns grouped so the main view stays readable
wa.column_dimensions.group("J", "T", hidden=False, outline_level=1)
wa.column_dimensions.group("AA", "AI", hidden=False, outline_level=1)
wa["U4"].comment = Comment("If the SKU is not short: allocation = need to target.\nIf short: priority floor first, then the rest in proportion to forecast (BR-02), solved so that all supply is used (depots capped at their need, the rest shared by forecast); rounded down to whole cases.", "Chiamaka Ikpo")

# =========================================================== Transfers
wt = wb.create_sheet("Transfers")
wt["A1"], wt["A1"].font = "Suggested inter-depot transfers (BR-03)", FT
wt["A2"], wt["A2"].font = ("A transfer is suggested only if the depot with the most excess (above its max band) can cover the depot with the largest "
                          "shortfall (below its min band), and the transfer arrives before that depot runs out.", FN)
head(wt, 4, ["SKU", "Largest excess (cases)", "From depot", "Largest shortfall (cases)", "To depot", "To-depot projected cover (days)",
             "Transfer lead time (days)", "Suggestion", "Transfer qty (cases)"])
for i, (s_name, *_rest) in enumerate(SKUS, 5):
    ex = f"Allocation!$AA${R0}:$AA${AEND}"
    sh = f"Allocation!$AB${R0}:$AB${AEND}"
    sk = f"Allocation!$A${R0}:$A${AEND}"
    fm = {
        "A": f"=Inputs!A{26 + i - 5}",
        "B": f"=_xlfn.MAXIFS({ex},{sk},A{i})",
        "C": f'=IF(B{i}>0,INDEX(Allocation!$B${R0}:$B${AEND},MATCH(A{i}&"|"&B{i},Allocation!$AC${R0}:$AC${AEND},0)),"–")',
        "D": f"=_xlfn.MAXIFS({sh},{sk},A{i})",
        "E": f'=IF(D{i}>0,INDEX(Allocation!$B${R0}:$B${AEND},MATCH(A{i}&"|"&D{i},Allocation!$AD${R0}:$AD${AEND},0)),"–")',
        "F": f'=IF(D{i}>0,INDEX(Allocation!$V${R0}:$V${AEND},MATCH(A{i}&"|"&D{i},Allocation!$AD${R0}:$AD${AEND},0)),"–")',
        "G": f"={TLT}",
        "H": (f'=IF(OR(B{i}=0,D{i}=0),"No transfer needed",IF(B{i}<D{i},"No: excess too small to cover shortfall",'
              f'IF(G{i}>F{i},"No: would arrive after stockout","Transfer")))'),
        "I": f'=IF(H{i}="Transfer",D{i},0)',
    }
    for c_letter, formula in fm.items():
        cell = wt[f"{c_letter}{i}"]
        cell.value, cell.font, cell.border = formula, F, BOX
        if c_letter in ("B", "D", "I"):
            cell.number_format = "#,##0"
        if c_letter == "F":
            cell.number_format = "0.0"
wt.conditional_formatting.add("H5:H10", CellIsRule(operator="equal", formula=['"Transfer"'], fill=PatternFill("solid", fgColor="E2F3EA"), font=Font(name=A, bold=True, color="1F7A4D")))
widths(wt, [28, 13, 18, 14, 18, 15, 13, 36, 13])
wt.row_dimensions[4].height = 42

# =========================================================== Summary
wsu = wb.create_sheet("Summary", 1)
wsu["A1"], wsu["A1"].font = "Deployment summary", FT
wsu["A2"], wsu["A2"].font = "Imbalance index = spread (standard deviation) of days of cover across depots. Lower = stock more evenly placed.", FN
head(wsu, 4, ["SKU", "Class", "Supply (cases)", "Total need", "Allocated", "Unallocated", "Short?", "Fill of need",
              "Imbalance before (days)", "Imbalance after (days)", "Exceptions before", "Exceptions after", "Transfer"])
sk = f"Allocation!$A${R0}:$A${AEND}"
for i in range(5, 11):
    src = 26 + i - 5
    G_ = f"Allocation!$G${R0}:$G${AEND}"
    V_ = f"Allocation!$V${R0}:$V${AEND}"
    fm = {
        "A": f"=Inputs!A{src}", "B": f"=Inputs!B{src}", "C": f"=Inputs!C{src}",
        "D": f"=SUMIF({sk},A{i},Allocation!$I${R0}:$I${AEND})",
        "E": f"=SUMIF({sk},A{i},Allocation!$U${R0}:$U${AEND})",
        "F": f"=MAX(0,C{i}-E{i})",
        "G": f'=IF(D{i}>C{i},"Yes","No")',
        "H": f"=IFERROR(E{i}/D{i},1)",
        "I": f"=SQRT(SUMPRODUCT(({sk}=A{i})*({G_}-AVERAGEIF({sk},A{i},{G_}))^2)/COUNTIF({sk},A{i}))",
        "J": f"=SQRT(SUMPRODUCT(({sk}=A{i})*({V_}-AVERAGEIF({sk},A{i},{V_}))^2)/COUNTIF({sk},A{i}))",
        "K": f'=COUNTIFS({sk},A{i},Allocation!$Z${R0}:$Z${AEND},"<>Within band")',
        "L": f'=COUNTIFS({sk},A{i},Allocation!$Y${R0}:$Y${AEND},"<>Within band")',
        "M": f"=INDEX(Transfers!$H$5:$H$10,MATCH(A{i},Transfers!$A$5:$A$10,0))",
    }
    for c_letter, formula in fm.items():
        cell = wsu[f"{c_letter}{i}"]
        cell.value, cell.font, cell.border = formula, F, BOX
        cell.number_format = {"H": "0%", "I": "0.0", "J": "0.0"}.get(c_letter, "#,##0")
wsu["A11"], wsu["A11"].font = "Total", FB
for c_letter in "CDEFKL":
    cell = wsu[f"{c_letter}11"]
    cell.value, cell.font, cell.border, cell.number_format = f"=SUM({c_letter}5:{c_letter}10)", FB, BOX, "#,##0"
wsu["H11"], wsu["H11"].font, wsu["H11"].number_format = "=IFERROR(E11/D11,1)", FB, "0%"
wsu["I11"], wsu["I11"].font, wsu["I11"].number_format = "=AVERAGE(I5:I10)", FB, "0.0"
wsu["J11"], wsu["J11"].font, wsu["J11"].number_format = "=AVERAGE(J5:J10)", FB, "0.0"

wsu["A14"], wsu["A14"].font = "KPIs for this cycle", FB
kpis = [
    ("Depot/SKU exceptions to review (of 30)", "=L11", "#,##0"),
    ("Share of rows planners must review (BR-04)", f"=L11/COUNTA(Allocation!$A${R0}:$A${AEND})", "0%"),
    ("Average imbalance index before → after (days)", '=TEXT(I11,"0.0")&" → "&TEXT(J11,"0.0")', "@"),
    ("SKUs short of supply", '=COUNTIF(G5:G10,"Yes")', "#,##0"),
    ("Transfers suggested", '=COUNTIF(M5:M10,"Transfer")', "#,##0"),
]
for i, (label, formula, fmt) in enumerate(kpis, 15):
    wsu.cell(row=i, column=1, value=label).font = F
    c = wsu.cell(row=i, column=3, value=formula)
    c.font, c.number_format, c.border = FB, fmt, BOX
widths(wsu, [28, 9, 12, 11, 11, 12, 8, 10, 13, 13, 12, 12, 30])
wsu.row_dimensions[4].height = 42
wsu.conditional_formatting.add("G5:G10", CellIsRule(operator="equal", formula=['"Yes"'], fill=PatternFill("solid", fgColor="FCE6E4"), font=Font(name=A, color="B42318", bold=True)))
wsu.conditional_formatting.add("M5:M10", CellIsRule(operator="equal", formula=['"Transfer"'], fill=PatternFill("solid", fgColor="E2F3EA"), font=Font(name=A, bold=True, color="1F7A4D")))

wb.save(OUT)
print("wrote", OUT.relative_to(ROOT))
